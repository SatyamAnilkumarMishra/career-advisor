"""Resume upload validation and text extraction.

Separate from `rag_pipeline.py` (which handles the *reference document* used
for RAG) because a resume isn't indexed into the vector store — it's read
once, turned into plain text, and handed to `career_tools.analyze_resume()`
for a single LLM pass. Supports PDF, DOCX, and plain text, since resumes are
commonly shared in all three.
"""

from __future__ import annotations

import logging
import os

from backend.config import Settings
from backend.errors import ResumeParsingError

logger = logging.getLogger(__name__)

_SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def validate_resume_file(file_path: str, *, max_size_bytes: int) -> str:
    """Validate the uploaded resume and return its lowercase extension."""
    if not os.path.exists(file_path):
        raise ResumeParsingError(f"File not found: {os.path.basename(file_path)}")

    size = os.path.getsize(file_path)
    if size == 0:
        raise ResumeParsingError("The uploaded resume is empty.")
    if size > max_size_bytes:
        max_mb = max_size_bytes / (1024 * 1024)
        raise ResumeParsingError(
            f"The uploaded resume is too large ({size / (1024 * 1024):.1f} MB). "
            f"Maximum allowed size is {max_mb:.0f} MB."
        )

    ext = os.path.splitext(file_path)[1].lower()
    if ext not in _SUPPORTED_EXTENSIONS:
        raise ResumeParsingError(
            f"Unsupported resume file type '{ext}'. Please upload a PDF, DOCX, or TXT file."
        )
    return ext


def _extract_pdf_text(file_path: str) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as exc:  # pragma: no cover
        raise ResumeParsingError(
            "Required PDF dependencies are not installed. Run: pip install -r requirements.txt"
        ) from exc

    try:
        reader = PdfReader(file_path)
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:
        raise ResumeParsingError(
            "We couldn't read this PDF resume. It may be corrupted, encrypted, or "
            "a scanned image without extractable text.",
            cause=exc,
        ) from exc
    return "\n".join(pages)


def _extract_docx_text(file_path: str) -> str:
    try:
        import docx
    except ImportError as exc:  # pragma: no cover
        raise ResumeParsingError(
            "Required DOCX dependencies are not installed. Run: pip install -r requirements.txt"
        ) from exc

    try:
        document = docx.Document(file_path)
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                paragraphs.extend(cell.text for cell in row.cells)
    except Exception as exc:
        raise ResumeParsingError(
            "We couldn't read this DOCX resume. It may be corrupted or in an "
            "unsupported format.",
            cause=exc,
        ) from exc
    return "\n".join(paragraphs)


def _extract_txt_text(file_path: str) -> str:
    try:
        with open(file_path, encoding="utf-8", errors="ignore") as f:
            return f.read()
    except Exception as exc:
        raise ResumeParsingError("We couldn't read this text file.", cause=exc) from exc


def extract_resume_text(
    file_path: str, settings: Settings | None = None, *, max_size_bytes: int | None = None
) -> str:
    """Validate and extract plain text from a resume file (PDF, DOCX, or TXT)."""
    if max_size_bytes is None:
        if settings is None:
            raise ValueError("Either `settings` or `max_size_bytes` must be provided.")
        max_size_bytes = settings.max_resume_size_bytes

    ext = validate_resume_file(file_path, max_size_bytes=max_size_bytes)

    logger.info("Extracting resume text from %s file: %s", ext, file_path)
    if ext == ".pdf":
        text = _extract_pdf_text(file_path)
    elif ext == ".docx":
        text = _extract_docx_text(file_path)
    else:
        text = _extract_txt_text(file_path)

    text = text.strip()
    if not text:
        raise ResumeParsingError(
            "No readable text was found in this resume. If it's a scanned "
            "image, try exporting a text-based PDF or pasting the text directly."
        )
    logger.info("Extracted %d character(s) of resume text", len(text))
    return text
